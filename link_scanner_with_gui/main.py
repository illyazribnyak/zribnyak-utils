import os
import re
import fitz
import openpyxl
import docx
import csv
from bs4 import BeautifulSoup
from pathlib import Path
from tkinter import Tk, Label, Button, filedialog, messagebox, StringVar
from html import escape
from urllib.parse import urlparse

SOCIAL_MEDIA = {
    "Facebook": ["facebook.com"],
    "Однокласники": ["ok.ru"],
    "Вконтакті": ["vk.com"],
    "X": ["twitter.com", "x.com"],
    "Threads": ["threads.net"],
    "LinkedIn": ["linkedin.com"],
    "TikTok": ["tiktok.com"],
    "YouTube": ["youtube.com"]
}


# Доповнено: Функція для перевірки, чи є це адреса пошти, яку треба ігнорувати
def is_email(url):
    return re.match(r"[^@]+@[^@]+\.[^@]+", url)


def extract_social_urls(text):
    urls = []
    pattern = r'https?://[^\s<>")]+|www\.[^\s<>")]+'
    matches = re.findall(pattern, text)

    for url in matches:
        if url.startswith("www."):
            url = "https://" + url
        if is_email(url):
            continue

        # Витягуємо домен
        parsed = urlparse(url)
        netloc = parsed.netloc.lower().replace("www.", "")

        for network, domains in SOCIAL_MEDIA.items():
            if netloc in domains:
                urls.append((network, url))
                break
    return urls



def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"[DOCX] {file_path}: {e}")
        return ""


def extract_text_from_pdf(file_path):
    try:
        text = ""
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text("text") + "\n"
                for link in page.get_links():
                    if "uri" in link:
                        text += link["uri"] + "\n"
        return text
    except Exception as e:
        print(f"[PDF] {file_path}: {e}")
        return ""


def extract_text_from_xlsx(file_path):
    try:
        text = ""
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                text += " ".join([str(cell.value) if cell.value else "" for cell in row]) + "\n"
        return text
    except Exception as e:
        print(f"[XLSX] {file_path}: {e}")
        return ""

def extract_text_from_txt(file_path):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        print(f"[TXT] {file_path}: {e}")
        return ""

def extract_text_from_csv(file_path):
    try:
        text = ""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                text += " ".join(row) + "\n"
        return text
    except Exception as e:
        print(f"[CSV] {file_path}: {e}")
        return ""

def extract_text_from_html(file_path):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f, "html.parser")
            return soup.get_text()
    except Exception as e:
        print(f"[HTML] {file_path}: {e}")
        return ""


def save_results(social_links, output_path, file_format="txt"):
    for network, urls in social_links.items():
        if not urls:
            continue
        if file_format == "txt":
            file_path = output_path / f"{network}.txt"
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("\n".join(urls))
        elif file_format == "csv":
            file_path = output_path / f"{network}.csv"
            with open(file_path, "w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["Network", "URL"])
                for url in urls:
                    writer.writerow([network, url])
        elif file_format == "html":
            file_path = output_path / f"{network}.html"
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("<html><body><h1>{}</h1><ul>".format(escape(network)))
                for url in urls:
                    f.write(f"<li><a href='{escape(url)}'>{escape(url)}</a></li>")
                f.write("</ul></body></html>")
        print(f"✅ Збережено {len(urls)} посилань у {file_path}")


def scan_folder_and_extract_urls(folder_path, output_path, file_format, status_label=None):
    output_path.mkdir(parents=True, exist_ok=True)
    social_links = {key: set() for key in SOCIAL_MEDIA.keys()}

    for root, _, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            ext = Path(file).suffix.lower()

            # ігнорувати результати попереднього сканування
            if file.startswith("Facebook.") or file.startswith("X.") or file.startswith("YouTube."):
                continue

            if ext == ".docx":
                text = extract_text_from_docx(file_path)
            elif ext == ".pdf":
                text = extract_text_from_pdf(file_path)
            elif ext == ".xlsx":
                text = extract_text_from_xlsx(file_path)
            elif ext == ".txt":
                text = extract_text_from_txt(file_path)
            elif ext == ".csv":
                text = extract_text_from_csv(file_path)
            elif ext in [".html", ".htm"]:
                text = extract_text_from_html(file_path)
            else:
                continue

            if text:
                urls = extract_social_urls(text)
                for network, url in urls:
                    social_links[network].add(url)

    save_results(social_links, output_path, file_format)
    if status_label:
        status_label.config(text="Готово!")
    messagebox.showinfo("Успіх", f"Сканування завершено!\nРезультати збережено у: {output_path}")


def create_app():
    root = Tk()
    root.title("Пошук соцмережевих посилань")
    root.geometry("600x300")
    root.resizable(False, False)

    scan_path = StringVar()
    save_path = StringVar()
    file_format = StringVar(value="txt")  # Значення за замовчуванням

    Label(root, text="Пошук соцмереж у файлах (.docx, .pdf, .xlsx, .txt, .csv, .html)", font=("Arial", 13)).pack(pady=15)

    status_label = Label(root, text="", fg="blue", font=("Arial", 10))
    status_label.pack()

    def select_scan_folder():
        folder = filedialog.askdirectory(title="Оберіть папку для сканування")
        if folder:
            scan_path.set(folder)
            status_label.config(text=f"Папка сканування: {folder}")

    def select_save_folder():
        folder = filedialog.askdirectory(title="Оберіть папку для збереження результатів")
        if folder:
            save_path.set(folder)
            status_label.config(text=f"Папка збереження: {folder}")

    def start_scan():
        if not scan_path.get() or not save_path.get():
            messagebox.showwarning("Увага", "Будь ласка, оберіть обидві папки перед запуском!")
            return
        scan_folder_and_extract_urls(Path(scan_path.get()), Path(save_path.get()), file_format.get(), status_label)

    def select_file_format():
        formats = [("TXT", "txt"), ("CSV", "csv"), ("HTML", "html")]
        file_format.set(filedialog.askstring("Вибір формату", "Оберіть формат файлів для збереження (txt, csv, html)",
                                             initialvalue="txt"))

    Button(root, text="1. Обрати папку сканування", font=("Arial", 11), command=select_scan_folder).pack(pady=5)
    Button(root, text="2. Обрати папку збереження", font=("Arial", 11), command=select_save_folder).pack(pady=5)
    Button(root, text="3. Почати сканування", font=("Arial", 11, "bold"), command=start_scan).pack(pady=10)
    Button(root, text="4. Вийти", font=("Arial", 10), command=root.destroy).pack(pady=5)

    root.mainloop()


if __name__ == "__main__":
    create_app()
